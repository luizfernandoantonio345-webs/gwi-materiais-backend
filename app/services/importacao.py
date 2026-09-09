import csv
import io
import unicodedata
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook

MAX_LINHAS = 5000


def _sem_acento(texto: str) -> str:
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def normalizar_chave(texto: str) -> str:
    return _sem_acento(str(texto or "")).strip().lower().replace(" ", "_").replace("-", "_")


def parse_valor_br(valor) -> Decimal | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    texto = str(valor).strip()
    texto = texto.replace("R$", "").replace("r$", "").replace(" ", "").replace("\xa0", "")
    if not texto:
        return None
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(texto).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def _ler_csv(conteudo: bytes) -> list[list[str]]:
    texto = None
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = conteudo.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        texto = conteudo.decode("latin-1", errors="replace")
    amostra = texto[:2048]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
        delim = dialeto.delimiter
    except csv.Error:
        delim = ";" if amostra.count(";") >= amostra.count(",") else ","
    leitor = csv.reader(io.StringIO(texto), delimiter=delim)
    return [list(linha) for linha in leitor]


def _ler_xlsx(conteudo: bytes) -> list[list[str]]:
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    ws = wb.active
    linhas: list[list] = []
    for linha in ws.iter_rows(values_only=True):
        linhas.append(["" if c is None else c for c in linha])
    wb.close()
    return linhas


def ler_planilha(conteudo: bytes, nome_arquivo: str) -> list[dict]:
    nome = (nome_arquivo or "").lower()
    if nome.endswith(".xlsx"):
        matriz = _ler_xlsx(conteudo)
    elif nome.endswith(".csv") or nome.endswith(".txt"):
        matriz = _ler_csv(conteudo)
    else:
        raise ValueError("Formato não suportado. Envie um arquivo .xlsx ou .csv.")

    matriz = [linha for linha in matriz if any(str(c).strip() for c in linha)]
    if not matriz:
        raise ValueError("A planilha está vazia.")
    if len(matriz) - 1 > MAX_LINHAS:
        raise ValueError(f"A planilha excede o limite de {MAX_LINHAS} linhas.")

    cabecalho = [normalizar_chave(c) for c in matriz[0]]
    registros: list[dict] = []
    for i, linha in enumerate(matriz[1:], start=2):
        dados = {}
        for idx, chave in enumerate(cabecalho):
            if not chave:
                continue
            valor = linha[idx] if idx < len(linha) else ""
            dados[chave] = str(valor).strip() if valor is not None else ""
        registros.append({"linha": i, "dados": dados})
    return registros


def valor_por_sinonimo(dados: dict, sinonimos: list[str]) -> str:
    for s in sinonimos:
        if s in dados and str(dados[s]).strip():
            return str(dados[s]).strip()
    return ""


def gerar_modelo_xlsx(titulo: str, colunas: list[str], exemplo: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    ws.append(colunas)
    if exemplo:
        ws.append(exemplo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gerar_relatorio_xlsx(linhas: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultado"
    ws.append(["Linha", "Ação", "Erro", "Dados"])
    for item in linhas:
        ws.append([
            item.get("linha"),
            item.get("acao"),
            item.get("erro") or "",
            ", ".join(f"{k}={v}" for k, v in (item.get("dados") or {}).items()),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
